"""הר הביטוח history, re-fetch question and limits — local dev DB, no browser, no model.

    source backend/venv/bin/activate && python backend/tests/test_harb_history.py

Fetch 1 = the real example export. Fetch 2 = a modified copy: one premium changed (Harel סיעודי
136.29 → 150), one policy gone (Clal 6939773), one new policy (מגדל 999888777).
"""
import asyncio
import io
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from sqlalchemy import delete, func, select, update  # noqa: E402

from app.config import settings  # noqa: E402
from app.database import async_session  # noqa: E402
from app.models import HarbRequest, InsurancePolicy, PolicyDocument, User, WorkerHeartbeat  # noqa: E402
from app.services.agent import registry, router  # noqa: E402
from app.services.agent.context import ToolContext  # noqa: E402
from app.services.policies import embeddings, harb_ingest, harb_jobs, store  # noqa: E402

XLSX = Path("/mnt/c/Users/roygi/Downloads/REPORT/police_example/עמיקם עינב הר הביטוח.xlsx")
IDN = "77777770"
TMP = Path("/tmp/claude-harb-history")
FAILS: list[str] = []


def check(cond, msg):
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        FAILS.append(msg)


def modified_copy() -> Path:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.worksheets[0]
    drop = None
    for row in ws.iter_rows(min_row=1):
        vals = [c.value for c in row]
        if vals[7] == 301611265:            # Harel סיעודי
            row[5].value = 150
        if vals[7] == 6939773:              # Clal life → gone
            drop = row[0].row
    ws.delete_rows(drop)
    last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 4).value and "בע\"מ" in str(ws.cell(r, 4).value))
    ws.insert_rows(last + 1)
    for col, v in enumerate(["ביטוח חיים", "ביטוח חיים למקרה מוות", "פוליסת ביטוח", 'מגדל חברה לביטוח בע"מ',
                             "01/01/2026 - 01/01/2040", 55.5, "חודשית", 999888777, "אישי"], start=1):
        ws.cell(last + 1, col).value = v
    TMP.mkdir(exist_ok=True)
    p = TMP / "harb_modified.xlsx"
    wb.save(p)
    return p


async def fresh(db, u):
    await db.execute(delete(PolicyDocument).where(PolicyDocument.user_id == u.id, PolicyDocument.customer_id_number == IDN))
    await db.execute(delete(InsurancePolicy).where(InsurancePolicy.user_id == u.id, InsurancePolicy.customer_id_number == IDN))
    await db.execute(delete(HarbRequest).where(HarbRequest.user_id == u.id, HarbRequest.customer_id_number == IDN))
    await db.execute(update(WorkerHeartbeat).where(WorkerHeartbeat.user_id == u.id).values(last_seen=datetime.utcnow()))
    await db.commit()


async def fetch(db, u, path: Path, ago_hours: float = 0) -> HarbRequest:
    req = HarbRequest(user_id=u.id, customer_id_number=IDN, customer_name="לקוח היסטוריה", birth_date=date(1980, 1, 1),
                      id_issue_date=date(2000, 1, 1), status="downloading",
                      created_at=datetime.utcnow() - timedelta(hours=ago_hours))
    db.add(req)
    await db.commit()
    await harb_ingest.ingest(db, req, path)
    if ago_hours:
        req.completed_at = datetime.utcnow() - timedelta(hours=ago_hours)
        await db.commit()
    return req


async def main():
    registry._load()
    async with async_session() as db:
        u = (await db.execute(select(User).where(User.email == "test@test.com"))).scalar_one()
        await fresh(db, u)
        print("history + what changed")
        r1 = await fetch(db, u, XLSX, ago_hours=48)        # two days ago (outside the 24h limits)
        check(r1.changes is None, "first fetch has no 'changes'")
        r2 = await fetch(db, u, modified_copy())
        ch = r2.changes or {}
        check(len(ch.get("new_policies", [])) == 1 and ch["new_policies"][0]["policy_number"] == "999888777", "new policy found (מגדל 999888777)")
        check(len(ch.get("removed_policies", [])) == 1 and ch["removed_policies"][0]["policy_number"] == "6939773", "removed policy found (כלל 6939773)")
        pc = ch.get("premium_changes", [])
        check(len(pc) == 1 and pc[0]["old"] == 136.29 and pc[0]["new"] == 150, f"one premium change 136.29 → 150 (got {pc})")
        check(not ch.get("period_changes"), "no false period changes")
        cur = (await db.execute(select(func.count()).select_from(InsurancePolicy).where(
            InsurancePolicy.user_id == u.id, InsurancePolicy.customer_id_number == IDN, InsurancePolicy.is_current.is_(True)))).scalar_one()
        old = (await db.execute(select(func.count()).select_from(InsurancePolicy).where(
            InsurancePolicy.user_id == u.id, InsurancePolicy.customer_id_number == IDN, InsurancePolicy.is_current.is_(False)))).scalar_one()
        check(cur == 50 and old == 50, f"current = the new fetch (50), the old fetch kept as history (50) — got {cur}/{old}")
        pic = await store.customer_picture(db, u.id, IDN)
        check(len(pic["history"]) == 2 and pic["history"][0]["current"] and not pic["history"][1]["current"], "history lists both fetches, newest current")
        check(pic["changes"] and pic["changes"]["summary_he"].startswith("1 פוליסות חדשות"), f"picture carries the changes ({pic['changes'] and pic['changes']['summary_he']})")
        check(any(p["policy_number"] == "999888777" for p in pic["policies"]) and not any(p["policy_number"] == "6939773" for p in pic["policies"]),
              "the picture shows ONLY the current fetch")
        old_doc = await db.get(PolicyDocument, __import__("uuid").UUID(pic["history"][1]["document_id"]))
        check(old_doc and not old_doc.is_current and "6939773" in old_doc.markdown, "the old portfolio document stays readable")
        new_doc = next(d for d in (await db.execute(select(PolicyDocument).where(
            PolicyDocument.user_id == u.id, PolicyDocument.customer_id_number == IDN, PolicyDocument.is_current.is_(True),
            PolicyDocument.source == "harb_portfolio"))).scalars())
        check("## שינויים מאז השליפה הקודמת" in new_doc.markdown and "999888777" in new_doc.markdown, "the new portfolio Markdown leads with what changed (AI can answer it)")
        hits = await embeddings.search_words(db, u.id, "6939773", customer_id=IDN)
        check(all("לא מופיעה יותר" in h["text"] for h in hits), "search never returns the superseded fetch as a live policy")
        from app.services import data_map
        m = await data_map.load(db, u)
        page = data_map.render(m, f"customers/{IDN}/policies.md")
        check("## שינויים מאז השליפה הקודמת" in page and "## שליפות קודמות" in page, "the data map page shows changes + previous fetches")

        print("ask before a re-fetch")
        ctx = ToolContext(db=db, user=u)
        args = {"id_number": IDN, "birth_date": "01/01/1980", "issue_date": "01/01/2000"}
        out = await registry.dispatch(ctx, "propose_harb_fetch", args)
        check(str(out).startswith("שאל את הסוכן: ") and "היום" in str(out) and not ctx.proposals,
              f"fetched today → a question, no card ({str(out)[:110]})")
        q = "תביא לי מהר הביטוח 77777770 01/01/1980 01/01/2000"
        check("confirm_refetch" not in router.harb_request(q), "a plain request is not a confirmation")
        check(router.harb_request(q + " שוב").get("confirm_refetch") is True, "'שוב' in the request = confirmed")
        print("limits")
        out = await registry.dispatch(ctx, "propose_harb_fetch", {**args, "confirm_refetch": True})
        check("פעמים ב-24" in str(out) or ctx.proposals, f"customer limit: 1 fetch in 24h (limit {settings.HARB_CUSTOMER_DAILY_LIMIT}) → still allowed")
        ctx.proposals.clear()
        settings.HARB_CUSTOMER_DAILY_LIMIT = 1
        out = await registry.dispatch(ctx, "propose_harb_fetch", {**args, "confirm_refetch": True})
        check("פעמים ב-24" in str(out) and not ctx.proposals, f"customer limit reached → refused ({str(out)[:90]})")
        settings.HARB_CUSTOMER_DAILY_LIMIT, settings.HARB_DAILY_LIMIT = 2, 1
        out = await registry.dispatch(ctx, "propose_harb_fetch", {**args, "id_number": "88888886", "confirm_refetch": True})
        check("מכסה" in str(out) and not ctx.proposals, f"daily limit reached → refused ({str(out)[:90]})")
        settings.HARB_DAILY_LIMIT, settings.HARB_MAX_QUEUE = 30, 0
        out = await registry.dispatch(ctx, "propose_harb_fetch", {**args, "id_number": "88888886", "confirm_refetch": True})
        check("בתור" in str(out) and not ctx.proposals, f"queue full → refused ({str(out)[:90]})")
        settings.HARB_MAX_QUEUE = 10
        # a fetch that died at login does not use up the quota
        await db.execute(update(HarbRequest).where(HarbRequest.id == r2.id).values(status="failed"))
        await db.commit()
        settings.HARB_CUSTOMER_DAILY_LIMIT = 1
        ctx.proposals.clear()
        out = await registry.dispatch(ctx, "propose_harb_fetch", {**args, "confirm_refetch": True})
        check(bool(ctx.proposals), f"a failed login doesn't count against the limit ({str(out)[:80]})")
        settings.HARB_CUSTOMER_DAILY_LIMIT = 2
        await fresh(db, u)
    print(f"\n{'ALL PASSED' if not FAILS else f'{len(FAILS)} FAILED'}")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    asyncio.run(main())
